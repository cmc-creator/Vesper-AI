"""
File Reader tool for Vesper — reads PDFs, DOCX, CSV, TXT, HTML from URL or Google Drive.

Actions:
  read_url    — fetch and extract text from any URL (PDF, DOCX, CSV, TXT, HTML)
  read_drive  — read a file from Google Drive by file ID (any supported format)
  summarize   — read + AI-summarize a file (pass content to ai_router)
  list_recent — list recently modified files in Google Drive
  search_drive — search Google Drive for files by name/query
"""

import os
import io
import re
import csv
import asyncio
import tempfile
import mimetypes
from typing import Optional

import aiohttp

# Optional heavy deps — fail gracefully
try:
    import pypdf
    _HAS_PYPDF = True
except ImportError:
    try:
        import PyPDF2 as pypdf
        _HAS_PYPDF = True
    except ImportError:
        _HAS_PYPDF = False

try:
    from docx import Document as DocxDocument
    _HAS_DOCX = True
except ImportError:
    _HAS_DOCX = False


def _strip_html(html: str) -> str:
    """Very simple HTML stripper - no BeautifulSoup dep."""
    html = re.sub(r'<style[^>]*>.*?</style>', '', html, flags=re.DOTALL | re.IGNORECASE)
    html = re.sub(r'<script[^>]*>.*?</script>', '', html, flags=re.DOTALL | re.IGNORECASE)
    html = re.sub(r'<[^>]+>', ' ', html)
    html = re.sub(r'&nbsp;', ' ', html)
    html = re.sub(r'&amp;', '&', html)
    html = re.sub(r'&lt;', '<', html)
    html = re.sub(r'&gt;', '>', html)
    html = re.sub(r'&quot;', '"', html)
    html = re.sub(r'\s{3,}', '\n\n', html)
    return html.strip()


def _extract_pdf_bytes(data: bytes) -> str:
    if not _HAS_PYPDF:
        return "[PDF extraction requires pypdf: pip install pypdf]"
    try:
        reader = pypdf.PdfReader(io.BytesIO(data))
        pages = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception:
                pass
        return "\n\n".join(pages)
    except Exception as e:
        return f"[PDF parse error: {e}]"


def _extract_docx_bytes(data: bytes) -> str:
    if not _HAS_DOCX:
        return "[DOCX extraction requires python-docx: pip install python-docx]"
    try:
        doc = DocxDocument(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as e:
        return f"[DOCX parse error: {e}]"


def _extract_csv_bytes(data: bytes) -> str:
    try:
        text = data.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        if not rows:
            return "(empty CSV)"
        # Format as markdown table
        lines = []
        if rows:
            header = rows[0]
            lines.append("| " + " | ".join(header) + " |")
            lines.append("| " + " | ".join(["---"] * len(header)) + " |")
            for row in rows[1:101]:  # max 100 rows shown
                lines.append("| " + " | ".join(str(c) for c in row) + " |")
            if len(rows) > 101:
                lines.append(f"\n... {len(rows) - 101} more rows (total {len(rows)} rows)")
        return "\n".join(lines)
    except Exception as e:
        return f"[CSV parse error: {e}]"


def _auto_extract(data: bytes, content_type: str, url: str = "") -> str:
    ct = content_type.lower()
    url_lower = url.lower()

    if "pdf" in ct or url_lower.endswith(".pdf"):
        return _extract_pdf_bytes(data)
    elif "spreadsheetml" in ct or "excel" in ct or url_lower.endswith((".xlsx", ".xls")):
        # Try openpyxl
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f"## Sheet: {ws.title}")
                for row in ws.iter_rows(max_row=100, values_only=True):
                    lines.append("| " + " | ".join(str(c) if c is not None else "" for c in row) + " |")
            return "\n".join(lines)
        except ImportError:
            return "[Excel extraction requires openpyxl: pip install openpyxl]"
        except Exception as e:
            return f"[Excel parse error: {e}]"
    elif "csv" in ct or url_lower.endswith(".csv"):
        return _extract_csv_bytes(data)
    elif "wordprocessingml" in ct or "msword" in ct or url_lower.endswith((".docx", ".doc")):
        return _extract_docx_bytes(data)
    elif "html" in ct or url_lower.endswith((".html", ".htm")):
        return _strip_html(data.decode("utf-8", errors="replace"))
    elif "text" in ct or url_lower.endswith((".txt", ".md", ".rst", ".py", ".js", ".ts", ".json")):
        return data.decode("utf-8", errors="replace")
    else:
        # Try UTF-8 text as last resort
        try:
            return data.decode("utf-8", errors="replace")
        except Exception:
            return "[Unable to extract text from this file type]"


async def _download_bytes(url: str, timeout: int = 30) -> tuple[bytes, str]:
    """Download bytes + content_type from a URL."""
    async with aiohttp.ClientSession() as s:
        async with s.get(url, timeout=aiohttp.ClientTimeout(total=timeout),
                         headers={"User-Agent": "Mozilla/5.0 (compatible; Vesper/1.0)"}) as r:
            content_type = r.headers.get("Content-Type", "application/octet-stream")
            data = await r.read()
            return data, content_type


def _truncate(text: str, max_chars: int = 12000) -> str:
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + f"\n\n... [truncated — {len(text)} total chars, showing first {max_chars}]"


async def file_reader_tool(params: dict, **kwargs) -> dict:
    action = params.get("action", "read_url").lower()
    max_chars = int(params.get("max_chars", 12000))

    if action == "read_url":
        url = params.get("url", "").strip()
        if not url:
            return {"error": "url is required for action=read_url"}
        try:
            data, content_type = await _download_bytes(url)
            text = _auto_extract(data, content_type, url)
            text = _truncate(text, max_chars)
            word_count = len(text.split())
            preview = text[:500] + ("..." if len(text) > 500 else "")
            return {
                "url": url,
                "content_type": content_type,
                "word_count": word_count,
                "char_count": len(text),
                "content": text,
                "preview": f"📄 Read {word_count} words from {url}\n\n{preview}",
            }
        except Exception as e:
            return {"error": f"Failed to read URL: {e}"}

    elif action == "read_drive":
        file_id = params.get("file_id", "").strip()
        if not file_id:
            return {"error": "file_id is required for action=read_drive"}
        try:
            from main import get_google_service
            drive = get_google_service("drive", "v3")

            # Get file metadata
            meta = drive.files().get(fileId=file_id, fields="name,mimeType,size").execute()
            name = meta.get("name", file_id)
            mime = meta.get("mimeType", "")

            # Google Workspace files — export them
            export_map = {
                "application/vnd.google-apps.document": ("text/plain", ".txt"),
                "application/vnd.google-apps.spreadsheet": ("text/csv", ".csv"),
                "application/vnd.google-apps.presentation": ("text/plain", ".txt"),
            }
            if mime in export_map:
                export_mime, ext = export_map[mime]
                req = drive.files().export_media(fileId=file_id, mimeType=export_mime)
                data = req.execute()
                if isinstance(data, str):
                    data = data.encode("utf-8")
            else:
                # Binary download
                req = drive.files().get_media(fileId=file_id)
                data = req.execute()
                if isinstance(data, str):
                    data = data.encode("utf-8")

            text = _auto_extract(data, mime, name)
            text = _truncate(text, max_chars)
            word_count = len(text.split())
            preview = text[:500] + ("..." if len(text) > 500 else "")
            return {
                "file_id": file_id,
                "name": name,
                "mime_type": mime,
                "word_count": word_count,
                "content": text,
                "preview": f"📄 Read '{name}' — {word_count} words\n\n{preview}",
            }
        except Exception as e:
            return {"error": f"Drive read error: {e}"}

    elif action == "summarize":
        # Read a file then summarize it with AI
        url = params.get("url", "").strip()
        file_id = params.get("file_id", "").strip()
        prompt_override = params.get("prompt", "")

        if url:
            read_result = await file_reader_tool({"action": "read_url", "url": url, "max_chars": max_chars})
        elif file_id:
            read_result = await file_reader_tool({"action": "read_drive", "file_id": file_id, "max_chars": max_chars})
        else:
            return {"error": "url or file_id required for action=summarize"}

        if "error" in read_result:
            return read_result

        content = read_result.get("content", "")
        name = read_result.get("name", read_result.get("url", "file"))

        ai_router = kwargs.get("ai_router")
        TaskType = kwargs.get("TaskType")
        if not ai_router or not TaskType:
            return {**read_result, "summary": "(ai_router not available — here is the raw content)"}

        prompt = prompt_override or f"Summarize this document clearly and thoroughly:\n\n{content[:10000]}"
        summary = await ai_router.chat(
            messages=[{"role": "user", "content": prompt}],
            task_type=TaskType.ANALYSIS,
            max_tokens=1500,
        )
        return {
            **read_result,
            "summary": summary,
            "preview": f"📄 **Summary of '{name}'**\n\n{summary}",
        }

    elif action == "list_recent":
        max_results = int(params.get("max_results", 20))
        folder_id = params.get("folder_id", "")
        try:
            from main import get_google_service
            drive = get_google_service("drive", "v3")
            query = "trashed=false"
            if folder_id:
                query += f" and '{folder_id}' in parents"
            result = drive.files().list(
                q=query,
                pageSize=max_results,
                orderBy="modifiedTime desc",
                fields="files(id,name,mimeType,modifiedTime,size,webViewLink)"
            ).execute()
            files = result.get("files", [])
            lines = [f"📁 **{len(files)} recent files**\n"]
            for f in files:
                icon = "📄" if "document" in f.get("mimeType", "") else \
                       "📊" if "spreadsheet" in f.get("mimeType", "") else \
                       "📑" if "presentation" in f.get("mimeType", "") else "📎"
                modified = f.get("modifiedTime", "")[:10]
                lines.append(f"{icon} **{f['name']}** — modified {modified}")
                lines.append(f"  ID: `{f['id']}` | [Open]({f.get('webViewLink', '')})")
            return {"files": files, "preview": "\n".join(lines)}
        except Exception as e:
            return {"error": f"Drive list error: {e}"}

    elif action == "search_drive":
        query = params.get("query", "").strip()
        if not query:
            return {"error": "query is required for action=search_drive"}
        max_results = int(params.get("max_results", 20))
        try:
            from main import get_google_service
            drive = get_google_service("drive", "v3")
            result = drive.files().list(
                q=f"name contains '{query}' and trashed=false",
                pageSize=max_results,
                fields="files(id,name,mimeType,modifiedTime,webViewLink)"
            ).execute()
            files = result.get("files", [])
            lines = [f"🔍 **Search results for '{query}'** — {len(files)} files\n"]
            for f in files:
                modified = f.get("modifiedTime", "")[:10]
                lines.append(f"📎 **{f['name']}** — {modified}")
                lines.append(f"  ID: `{f['id']}` | [Open]({f.get('webViewLink', '')})")
            return {"files": files, "preview": "\n".join(lines)}
        except Exception as e:
            return {"error": f"Drive search error: {e}"}

    else:
        return {"error": f"Unknown action '{action}'. Use: read_url | read_drive | summarize | list_recent | search_drive"}
